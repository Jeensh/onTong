"""OD-11-D2-1 테스트용 fixture helpers.

pypdf / python-docx / python-pptx 로 최소 샘플 파일을 생성한다.
테스트 단위에서 `tmp_path / "..."` 에 저장해 실제 파싱을 검증.
"""
from __future__ import annotations

from pathlib import Path


# ---------------------------------------------------------------------------
# PDF : pypdf 만으로 content stream 을 직접 작성해 텍스트 삽입
# ---------------------------------------------------------------------------
def make_pdf(path: Path, pages_text: list[str]) -> None:
    """주어진 텍스트 리스트로 단순 PDF 생성. 각 entry = 1 page.

    폰트는 Helvetica 고정, Y=720 위치 고정. 텍스트 안의 `(` `)` 는 이스케이프.
    """
    from pypdf import PdfWriter
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
        NumberObject,
    )

    w = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = w._add_object(font)

    for text in pages_text:
        page = w.add_blank_page(width=612, height=792)
        escaped = text.replace("\\", "\\\\").replace("(", r"\(").replace(")", r"\)")
        cs_bytes = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1", errors="replace")
        stream = DecodedStreamObject()
        stream.set_data(cs_bytes)
        stream_ref = w._add_object(stream)
        page[NameObject("/Contents")] = stream_ref
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        })
        # MediaBox 이미 blank page 에 설정됨
        _ = ArrayObject, NumberObject   # silence unused warnings

    with path.open("wb") as f:
        w.write(f)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def make_docx(path: Path, *, title: str, headings: list[tuple[int, str, str]]) -> None:
    """python-docx 로 간단 DOCX 생성.

    headings = [(level, heading_text, body_text), ...]
    title 은 문서 최상단 Heading 1.
    """
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for level, heading, body in headings:
        doc.add_heading(heading, level=level)
        if body:
            doc.add_paragraph(body)
    doc.save(str(path))


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def make_pptx(path: Path, slides: list[tuple[str, str]]) -> None:
    """python-pptx 로 간단 PPTX 생성.

    slides = [(title, body_text), ...]
    """
    from pptx import Presentation

    prs = Presentation()
    title_body_layout = prs.slide_layouts[1]   # "Title and Content"
    for title, body in slides:
        slide = prs.slides.add_slide(title_body_layout)
        slide.shapes.title.text = title
        # body placeholder
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = body
    prs.save(str(path))
